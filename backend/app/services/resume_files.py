"""Resume file checks and text extraction (PDF via pypdf, DOCX via python-docx)."""

import io
import zipfile
from dataclasses import dataclass
from pathlib import PurePath

from app.core.errors import AppError, ValidationAppError

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# Less text than this means a scanned/image-only resume we cannot read.
MIN_TEXT_CHARS = 200


class UnsupportedFileError(AppError):
    status_code = 415
    code = "UNSUPPORTED_FILE_TYPE"
    message = "Only PDF and DOCX resumes are supported."


class FileTooLargeError(AppError):
    status_code = 413
    code = "FILE_TOO_LARGE"
    message = "The file is too large."


@dataclass(frozen=True)
class ResumeFile:
    kind: str  # "pdf" | "docx"
    mime_type: str


def detect(filename: str, data: bytes) -> ResumeFile:
    """Trust neither the extension nor the browser's content type alone: check the bytes."""
    extension = PurePath(filename).suffix.lower()
    if extension == ".pdf" and data.startswith(b"%PDF"):
        return ResumeFile("pdf", PDF_MIME)
    if extension == ".docx" and data.startswith(b"PK") and _is_docx(data):
        return ResumeFile("docx", DOCX_MIME)
    raise UnsupportedFileError()


def _is_docx(data: bytes) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            return "word/document.xml" in archive.namelist()
    except zipfile.BadZipFile:
        return False


def extract_text(data: bytes, kind: str) -> str:
    """Plain text of the resume. Raises a 422 when nothing readable is inside."""
    try:
        text = _pdf_text(data) if kind == "pdf" else _docx_text(data)
    except Exception as exc:
        raise ValidationAppError(
            "The file could not be read. Is it damaged or password-protected?",
            code="RESUME_UNREADABLE",
        ) from exc
    text = _normalize(text)
    if len(text) < MIN_TEXT_CHARS:
        raise ValidationAppError(
            "Too little text found. Scanned (image-only) resumes are not supported.",
            code="RESUME_NO_TEXT",
        )
    return text


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise ValueError("encrypted PDF")
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _docx_text(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells: list[str] = []
            for cell in row.cells:
                if cell.text not in cells:  # merged cells repeat their text
                    cells.append(cell.text)
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def _normalize(text: str) -> str:
    lines = [" ".join(line.split()) for line in text.replace("\x00", "").splitlines()]
    out: list[str] = []
    for line in lines:
        if line or (out and out[-1]):  # collapse runs of blank lines
            out.append(line)
    return "\n".join(out).strip()
